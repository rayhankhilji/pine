"""Programmatic fixtures — tiny real documents built at test time (no committed binaries)."""

import io

import pymupdf
import pytest


def make_pdf_with_table() -> bytes:
    """One-page PDF with a 5x4 ruled table containing known cell text."""
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_text((72, 60), "Quarterly Revenue Report")
    headers = ["Metric", "Q1", "Q2", "Q3"]
    data = [
        ["Revenue", "100", "120", "140"],
        ["COGS", "40", "45", "50"],
        ["Gross profit", "60", "75", "90"],
        ["EBITDA", "10", "20", "30"],
    ]
    x0, y0, cw, ch = 72, 100, 120, 24
    rows = [headers] + data
    for r, row in enumerate(rows):
        for c, cell in enumerate(row):
            rect = pymupdf.Rect(
                x0 + c * cw, y0 + r * ch, x0 + (c + 1) * cw, y0 + (r + 1) * ch
            )
            page.draw_rect(rect)
            page.insert_text((rect.x0 + 4, rect.y0 + 16), cell, fontsize=9)
    return bytes(doc.tobytes())


def make_scanned_pdf() -> bytes:
    """PDF whose only page is a full-page raster image of text (no text layer)."""
    src = pymupdf.open()
    sp = src.new_page(width=612, height=792)
    sp.insert_textbox(
        pymupdf.Rect(72, 72, 540, 400),
        "Balance confirmation as of 31 December 2025: USD 6,200,000.00",
        fontsize=14,
    )
    pix = sp.get_pixmap(dpi=150)
    img = pix.tobytes("png")
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(page.rect, stream=img)
    return bytes(doc.tobytes())


def make_xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "P&L"
    ws.append(["Metric", "FY2024", "FY2025"])
    ws.append(["Revenue", 4100000, 9700000])
    ws.append(["COGS", 1230000, 2522000])
    ws2 = wb.create_sheet("Balance")
    ws2.append(["Item", "31-Dec-2025"])
    ws2.append(["Cash", 6200000])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Service Agreement", level=1)
    d.add_paragraph("This agreement is between the parties.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    pres = Presentation()
    slide = pres.slides.add_slide(pres.slide_layouts[0])
    slide.shapes.title.text = "Northwind Series B"
    body = pres.slides.add_slide(pres.slide_layouts[1])
    body.shapes.title.text = "Metrics"
    tf = body.placeholders[1].text_frame
    tf.text = "ARR reached $12.0M in Q4 2025"
    _ = Inches(1)  # keep import used
    buf = io.BytesIO()
    pres.save(buf)
    return buf.getvalue()


def make_eml(with_attachment: bool = True) -> bytes:
    from email.message import EmailMessage

    msg = EmailMessage()
    msg["From"] = "cfo@northwind.example"
    msg["To"] = "analyst@fund.example"
    msg["Date"] = "Tue, 20 Jan 2026 10:00:00 +0000"
    msg["Subject"] = "Financials attached"
    msg.set_content("Please find the summary attached.\nBest,\nCFO")
    if with_attachment:
        msg.add_attachment(
            make_pdf_with_table(),
            maintype="application",
            subtype="pdf",
            filename="02_Financials_summary.pdf",
        )
    return msg.as_bytes()


@pytest.fixture
def pdf_with_table() -> bytes:
    return make_pdf_with_table()


@pytest.fixture
def scanned_pdf() -> bytes:
    return make_scanned_pdf()


@pytest.fixture
def xlsx_two_sheets() -> bytes:
    return make_xlsx()


@pytest.fixture
def docx_file() -> bytes:
    return make_docx()


@pytest.fixture
def pptx_file() -> bytes:
    return make_pptx()


@pytest.fixture
def eml_with_attachment() -> bytes:
    return make_eml()
