"""T-F14-AC1: demo room builds deterministically; ground_truth.json validates."""

import csv
import io
import json
from pathlib import Path

import pytest

from pine.demo.build import build_demo
from pine.demo.schema import GroundTruth

EXPECTED_FILES = [
    "01_Northwind_SeriesB_Deck.pptx",
    "02_Financials_FY2024_FY2025.xlsx",
    "03_Customer_List.csv",
    "04_Bank_Statement_2025.csv",
    "05_Contracts/MSA_Acme_Corporation.pdf",
    "05_Contracts/MSA_Borealis_Health.pdf",
    "05_Contracts/SOW_Helios_Logistics.pdf",
    "05_Contracts/MSA_Delta_Freight.pdf",
    "05_Contracts/Order_Form_Kestrel_Bank.pdf",
    "06_Cap_Table.xlsx",
    "07_Board_Deck_Q4_2025.pdf",
    "08_Emails/email_01.eml",
    "08_Emails/email_02.eml",
    "08_Emails/email_03.eml",
    "08_Emails/email_04.eml",
    "08_Emails/email_05.eml",
    "08_Emails/email_06.eml",
    "09_Scanned_Bank_Letter.pdf",
    "ground_truth.json",
]


@pytest.fixture(scope="module")
def room(tmp_path_factory: pytest.TempPathFactory) -> Path:
    return build_demo(tmp_path_factory.mktemp("northwind"))


def test_all_files_exist_and_small(room: Path) -> None:
    for rel in EXPECTED_FILES:
        assert (room / rel).is_file(), rel
    total = sum(p.stat().st_size for p in room.rglob("*") if p.is_file())
    assert total < 5 * 1024 * 1024


def test_ground_truth_validates(room: Path) -> None:
    gt = GroundTruth.model_validate(json.loads((room / "ground_truth.json").read_text()))
    metrics = {(f.metric, f.value) for f in gt.facts}
    assert ("arr", 12_000_000) in metrics
    assert ("revenue", 9_700_000) in metrics
    assert ("cash_balance", 6_200_000) in metrics
    assert ("headcount", 84) in metrics
    assert ("tam", 40_000_000_000) in metrics
    rules = {c.rule_id for c in gt.contradictions}
    assert rules == {"R1", "R5"}
    r1 = next(c for c in gt.contradictions if c.rule_id == "R1")
    assert r1.expected_values == [12_000_000, 9_700_000, 10_200_000]
    assert r1.min_severity == "high"
    for f in gt.facts:
        assert (room / f.source_file).is_file(), f.source_file


def test_customer_list_sums(room: Path) -> None:
    lines = (room / "03_Customer_List.csv").read_text().splitlines()
    rows = list(csv.DictReader(lines))
    assert len(rows) == 40
    assert sum(int(r["annual_recurring_revenue"]) for r in rows) == 10_200_000
    top = max(rows, key=lambda r: int(r["annual_recurring_revenue"]))
    assert top["customer_name"] == "Acme Corporation"
    assert int(top["annual_recurring_revenue"]) == 2_244_000
    assert any(r["customer_name"] == "Helios Logistics" for r in rows)


def test_bank_statement(room: Path) -> None:
    text = (room / "04_Bank_Statement_2025.csv").read_text()
    assert text.startswith("ACCT 4839201156\n")
    rows = list(csv.DictReader(io.StringIO(text.split("\n", 1)[1])))
    credits = sum(int(r["credit"]) for r in rows if r["credit"])
    assert credits == 9_900_000
    assert int(rows[-1]["balance"]) == 6_200_000


def test_deck_content(room: Path) -> None:
    from pptx import Presentation

    prs = Presentation(str(room / "01_Northwind_SeriesB_Deck.pptx"))
    assert len(list(prs.slides)) == 12
    texts = [
        "\n".join(
            shape.text_frame.text for shape in slide.shapes if shape.has_text_frame
        )
        for slide in prs.slides
    ]
    assert "ARR reached $12.0M in Q4 2025, up 3x YoY" in texts[6]
    assert "TAM $40B" in texts[8]
    assert "$150M pre-money" in texts[10]


def test_cap_table_sums_103(room: Path) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(room / "06_Cap_Table.xlsx")
    ws = wb["Cap Table"]
    rows = list(ws.iter_rows(values_only=True))
    total = sum(float(r[3]) for r in rows[1:])
    assert abs(total - 103.0) < 0.001


def test_board_deck_pages(room: Path) -> None:
    import pymupdf

    doc = pymupdf.open(room / "07_Board_Deck_Q4_2025.pdf")
    assert doc.page_count == 8
    assert "Net burn $450k/month" in doc[2].get_text()
    assert "Runway 14 months" in doc[3].get_text()
    assert "NRR 118%" in doc[4].get_text()
    doc.close()


def test_contract_contents(room: Path) -> None:
    import pymupdf

    def page_text(path: Path) -> str:
        doc = pymupdf.open(path)
        text = " ".join(doc[0].get_text().split())
        doc.close()
        return text

    text = page_text(room / "05_Contracts/MSA_Acme_Corporation.pdf")
    assert "ACME Corporation" in text
    assert "Section 7.2" in text
    assert "$6.7M" in text

    ktext = page_text(room / "05_Contracts/Order_Form_Kestrel_Bank.pdf")
    assert "xclusiv" in ktext


def test_emails(room: Path) -> None:
    import email
    import email.policy

    e4 = email.message_from_bytes(
        (room / "08_Emails/email_04.eml").read_bytes(), policy=email.policy.default
    )
    atts = [p.get_filename() for p in e4.iter_attachments()]
    assert "Northwind_FY2025_Summary.pdf" in atts

    e6 = email.message_from_bytes(
        (room / "08_Emails/email_06.eml").read_bytes(), policy=email.policy.default
    )
    body = e6.get_body(("plain",)).get_content()  # type: ignore[union-attr]
    assert "Helios Logistics gave notice, churn effective Jan 2026" in body

    e5 = email.message_from_bytes(
        (room / "08_Emails/email_05.eml").read_bytes(), policy=email.policy.default
    )
    body5 = e5.get_body(("plain",)).get_content()  # type: ignore[union-attr]
    assert "pending litigation" in body5
    assert "$300k" in body5


def test_scanned_letter_has_no_text_layer(room: Path) -> None:
    import pymupdf

    doc = pymupdf.open(room / "09_Scanned_Bank_Letter.pdf")
    assert len(doc[0].get_text().strip()) < 10
    assert doc[0].get_images()
    doc.close()


def test_rebuild_requires_force(room: Path) -> None:
    with pytest.raises(FileExistsError):
        build_demo(room)
    build_demo(room, force=True)
