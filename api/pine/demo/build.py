"""Deterministic synthetic data-room generator (PRD F-14).

`build_demo(out_dir)` writes the Northwind SaaS Series B room plus
`ground_truth.json`. Seeded RNG (42) and fixed metadata timestamps keep output
stable across runs.
"""

import csv
import io
import json
import random
import shutil
from datetime import date, datetime
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path

from pine.demo.schema import (
    GroundTruth,
    GroundTruthContradiction,
    GroundTruthEntity,
    GroundTruthFact,
    Period,
)

FIXED_TS = datetime(2026, 1, 1)

# ---------------------------------------------------------------------------
# helpers


def _pdf(pages: list[list[str]], title: str = "") -> bytes:
    """Render one text line-block per page via PyMuPDF (no reportlab)."""
    import pymupdf

    doc = pymupdf.open()
    for lines in pages:
        page = doc.new_page(width=612, height=792)
        y = 72.0
        for line in lines:
            height = 30 + 14 * (len(line) // 80 + 1)
            page.insert_textbox(
                pymupdf.Rect(72, y, 540, y + height), line, fontsize=11
            )
            y += height + 6
    doc.set_metadata(
        {
            "producer": "pine-demo",
            "creator": "pine-demo",
            "title": title,
            "creationDate": "D:20260101000000Z",
            "modDate": "D:20260101000000Z",
        }
    )
    data = bytes(doc.tobytes(deflate=True))
    doc.close()
    return data


def _xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    for i, (name, rows) in enumerate(sheets.items()):
        ws = wb.active if i == 0 else wb.create_sheet(name)
        ws.title = name
        for row in rows:
            ws.append(row)
    wb.properties.created = FIXED_TS
    wb.properties.modified = FIXED_TS
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _eml(
    frm: str,
    to: str,
    when: datetime,
    subject: str,
    body: str,
    attachment: tuple[str, bytes] | None = None,
) -> bytes:
    msg = EmailMessage()
    msg["From"] = frm
    msg["To"] = to
    msg["Date"] = format_datetime(when)
    msg["Subject"] = subject
    msg.set_content(body)
    if attachment:
        msg.add_attachment(
            attachment[1],
            maintype="application",
            subtype="pdf",
            filename=attachment[0],
        )
    return msg.as_bytes()


# ---------------------------------------------------------------------------
# builders


def _deck_pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    title_body = prs.slide_layouts[1]

    def slide(title: str, bullets: list[str]) -> None:
        s = prs.slides.add_slide(title_body)
        s.shapes.title.text = title
        tf = s.placeholders[1].text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b

    slide(
        "Northwind — Series B",
        ["Private markets intelligence for SaaS", "January 2026"],
    )
    slide(
        "Problem",
        ["Investors reconcile data rooms by hand", "Weeks lost to bad metrics"],
    )
    slide("Product", ["Ingests the whole room", "Every claim backed by evidence"])
    slide("Why now", ["LLMs read documents well", "But need evidence"])
    slide("Traction", ["40 customers across NA/EMEA/APAC", "NRR 118%"])
    slide("Growth", ["ARR 3x year over year", "Land-and-expand motion"])
    slide("ARR", ["ARR reached $12.0M in Q4 2025, up 3x YoY"])
    slide("Competition", ["Generic RAG tools cannot reconcile numbers", "Pine is evidence-first"])
    slide("Market", ["TAM $40B", "Private markets diligence spend"])
    slide("Team", ["CEO — ex-Stripe", "CTO — ex-Palantir", "84 employees"])
    slide("Use of funds", ["$25M round at $150M pre-money", "50% engineering, 30% GTM, 20% ops"])
    slide("Contact", ["ir@northwind.example"])

    prs.core_properties.created = FIXED_TS
    prs.core_properties.modified = FIXED_TS
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _financials_xlsx() -> bytes:
    return _xlsx(
        {
            "P&L": [
                ["Metric", "FY2024", "FY2025"],
                ["SaaS revenue", 4100000, 9700000],
                ["COGS", 1230000, 2522000],
                ["Gross profit", 2870000, 7178000],
                ["Sales & marketing", 1900000, 3400000],
                ["R&D", 1600000, 2800000],
                ["G&A", 700000, 1100000],
                ["EBITDA", -1330000, -122000],
            ],
            "Balance": [
                ["Item", "2025-12-31"],
                ["Cash", 6200000],
                ["Accounts receivable", 900000],
                ["Total assets", 7100000],
                ["Deferred revenue", 1100000],
                ["Total liabilities", 1400000],
                ["Equity", 5700000],
            ],
            "Headcount": [
                ["Metric", "FY2024", "FY2025"],
                ["Headcount", 61, 84],
            ],
        }
    )


_FIRST = [
    "Atlas", "Beacon", "Cinder", "Drift", "Ember", "Fjord", "Gale", "Harbor",
    "Ionic", "Juniper", "Kelp", "Lumen", "Mesa", "Nimbus", "Orchid", "Prairie",
    "Quay", "Ridgeline", "Summit", "Tidal", "Umbra", "Vertex", "Willow",
    "Xenon", "Yonder", "Zephyr", "Alder", "Birch", "Cove", "Dune",
    "Estuary", "Flint", "Grove", "Hollow", "Isle",
]
_LAST = ["Systems", "Labs", "Group", "Works", "Analytics", "Freight", "Health",
         "Capital", "Robotics", "Media"]
_SEGMENTS = ["smb", "mid", "enterprise"]
_REGIONS = ["NA", "EMEA", "APAC"]


def _customer_list_csv() -> bytes:
    rng = random.Random(42)
    names = {"Acme Corporation", "Helios Logistics", "Borealis Health",
             "Delta Freight", "Kestrel Bank"}
    while len(names) < 40:
        names.add(f"{rng.choice(_FIRST)} {rng.choice(_LAST)}")
    others = sorted(names - {"Acme Corporation"})

    # distribute exactly 7,956,000 over 39 customers so the column sums to
    # exactly 10,200,000 with Acme at 2,244,000 (22% concentration)
    weights = [rng.randint(1, 100) for _ in others]
    total_w = sum(weights)
    amounts = [w * 7_956_000 // total_w for w in weights]
    amounts[-1] += 7_956_000 - sum(amounts)

    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(
        [
            "customer_name",
            "contract_start",
            "contract_end",
            "annual_recurring_revenue",
            "segment",
            "region",
        ]
    )

    def row(name: str, arr: int) -> list[object]:
        start = date(2024, rng.randint(1, 12), rng.randint(1, 28))
        end = date(start.year + 1, start.month, start.day)
        return [
            name,
            start.isoformat(),
            end.isoformat(),
            arr,
            rng.choice(_SEGMENTS),
            rng.choice(_REGIONS),
        ]

    writer.writerow(row("Acme Corporation", 2_244_000))
    for name, arr in zip(others, amounts, strict=True):
        writer.writerow(row(name, arr))
    return buf.getvalue().encode()


def _bank_statement_csv() -> bytes:
    rng = random.Random(43)
    # credits summing to exactly 9,900,000 across 12 months
    weights = [rng.randint(80, 120) for _ in range(12)]
    total_w = sum(weights)
    credits = [w * 9_900_000 // total_w for w in weights]
    credits[-1] += 9_900_000 - sum(credits)
    # debits summing to exactly 3,700,000 -> ending balance 6,200,000
    dweights = [rng.randint(80, 120) for _ in range(12)]
    dtotal = sum(dweights)
    debits = [w * 3_700_000 // dtotal for w in dweights]
    debits[-1] += 3_700_000 - sum(debits)

    buf = io.StringIO()
    buf.write("ACCT 4839201156\n")
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(["date", "description", "debit", "credit", "balance"])
    balance = 0
    for i in range(12):
        month = i + 1
        credit = credits[i]
        debit = debits[i]
        balance += credit
        writer.writerow(
            [f"2025-{month:02d}-15", "Customer receipts", "", credit, balance]
        )
        balance -= debit
        writer.writerow(
            [f"2025-{month:02d}-28", "Operating disbursements", debit, "", balance]
        )
    assert balance == 6_200_000
    return buf.getvalue().encode()


def _contracts() -> dict[str, bytes]:
    acme = _pdf(
        [
            [
                "MASTER SERVICES AGREEMENT",
                "This Master Services Agreement is entered into between Northwind, Inc. "
                "and ACME Corporation (the \"Customer\").",
                "Total contract value (TCV): $6.7M over a three (3) year term "
                "commencing January 1, 2025.",
                "Section 7.2 — Termination for Change of Control: Customer may "
                "terminate this Agreement upon a change of control of Northwind "
                "with thirty (30) days written notice.",
                "Governing law: State of Delaware.",
            ]
        ],
        "MSA — Acme Corporation",
    )
    borealis = _pdf(
        [
            [
                "MASTER SERVICES AGREEMENT",
                "Between Northwind, Inc. and Borealis Health.",
                "Term: twelve (12) months, renewing automatically for successive "
                "one-year terms unless either party gives sixty (60) days notice "
                "(auto-renewal).",
                "Annual fee: $312,000.",
            ]
        ],
        "MSA — Borealis Health",
    )
    helios = _pdf(
        [
            [
                "STATEMENT OF WORK — SOW-2025-014",
                "Between Northwind, Inc. and Helios Logistics under the Master "
                "Services Agreement dated March 2024.",
                "Scope: logistics analytics module. Fees: $480,000 annually.",
                "Term: April 2025 through March 2026.",
            ]
        ],
        "SOW — Helios Logistics",
    )
    delta = _pdf(
        [
            [
                "MASTER SERVICES AGREEMENT",
                "Between Northwind, Inc. and Delta Freight.",
                "Annual subscription fee of $260,000, invoiced quarterly.",
                "Standard limitation of liability and indemnification terms apply.",
            ]
        ],
        "MSA — Delta Freight",
    )
    kestrel = _pdf(
        [
            [
                "ORDER FORM — Kestrel Bank",
                "Products: Pine Analytics Suite, 40 seats.",
                "Exclusivity: during the term, Kestrel Bank shall not license any "
                "competing diligence platform (exclusivity clause, Section 3).",
                "Order value: $150,000 per year.",
            ]
        ],
        "Order Form — Kestrel Bank",
    )
    return {
        "MSA_Acme_Corporation.pdf": acme,
        "MSA_Borealis_Health.pdf": borealis,
        "SOW_Helios_Logistics.pdf": helios,
        "MSA_Delta_Freight.pdf": delta,
        "Order_Form_Kestrel_Bank.pdf": kestrel,
    }


def _cap_table_xlsx() -> bytes:
    # fully_diluted_pct deliberately sums to 103.0 (planted R5 contradiction)
    return _xlsx(
        {
            "Cap Table": [
                ["shareholder", "security_class", "shares", "fully_diluted_pct"],
                ["Founders", "common", 4_500_000, 45.0],
                ["Sequoia Scout Fund", "preferred_seed", 1_200_000, 12.0],
                ["Meridian Ventures", "preferred_series_a", 2_500_000, 25.0],
                ["ESOP", "options", 1_300_000, 13.0],
                ["Angel investors", "preferred_seed", 800_000, 8.0],
            ]
        }
    )


def _board_deck_pdf() -> bytes:
    pages = [
        ["Northwind — Q4 2025 Board Deck", "Prepared January 2026"],
        ["Agenda", "Metrics review", "Cash position", "Hiring plan", "Risks"],
        [
            "Burn",
            "Net burn $450k/month in Q4 2025.",
            "Gross burn $610k/month, collections $160k/month.",
        ],
        [
            "Runway",
            "Runway 14 months at current net burn.",
            "Cash on hand $6.2M as of December 31, 2025.",
        ],
        [
            "Retention",
            "NRR 118% for FY2025.",
            "Logo churn 3 accounts.",
        ],
        ["Hiring", "Headcount 84, plan to reach 110 by Q4 2026."],
        ["Risks", "Customer concentration; litigation exposure; churn notices."],
        ["Appendix", "Detailed financials in the data room."],
    ]
    return _pdf(pages, "Board Deck Q4 2025")


def _emails() -> dict[str, bytes]:
    cfo = "cfo@northwind.example"
    analyst = "analyst@fund.example"
    base = datetime(2026, 1, 12, 9, 0)

    summary_pdf = _pdf(
        [
            [
                "Northwind financial summary",
                "FY2025 SaaS revenue $9.7M; COGS $2.5M; gross profit $7.2M.",
                "Cash at 31 December 2025: $6.2M.",
            ]
        ],
        "Financial summary",
    )

    return {
        "email_01.eml": _eml(
            cfo, analyst, base, "Data room access",
            "Hi,\n\nSharing the Northwind data room. The deck, financials and "
            "contracts are all inside.\n\nBest,\nCFO",
        ),
        "email_02.eml": _eml(
            cfo, analyst, base.replace(hour=10), "Customer list update",
            "Attached earlier is the current customer list — 40 logos, "
            "$10.2M recurring revenue.\n\nCFO",
        ),
        "email_03.eml": _eml(
            "counsel@northwind.example", analyst, base.replace(day=13),
            "Contract notes",
            "Note the Acme MSA change-of-control clause (Section 7.2) and the "
            "Kestrel exclusivity clause when you review the contracts folder.",
        ),
        "email_04.eml": _eml(
            cfo, analyst, base.replace(day=14), "FY2025 financial summary",
            "Attached is a one-page summary of the FY2025 financials.\n\nCFO",
            attachment=("Northwind_FY2025_Summary.pdf", summary_pdf),
        ),
        "email_05.eml": _eml(
            "counsel@northwind.example", analyst, base.replace(day=15),
            "Pending litigation",
            "Flagging pending litigation with former contractor, exposure "
            "~$300k. Discovery is expected in Q1.\n\nCounsel",
        ),
        "email_06.eml": _eml(
            cfo, analyst, base.replace(day=16), "Helios notice",
            "Heads up — Helios Logistics gave notice, churn effective Jan 2026. "
            "We are trying to save the account.\n\nCFO",
        ),
    }


def _scanned_letter() -> bytes:
    """Image-only PDF: text rendered to PNG, inserted as an image."""
    import pymupdf

    src = pymupdf.open()
    sp = src.new_page(width=612, height=792)
    sp.insert_textbox(
        pymupdf.Rect(72, 72, 540, 400),
        "First Meridian Bank\n\nBalance confirmation as of 31 December 2025: "
        "USD 6,200,000.00\n\nAccount: 4839201156",
        fontsize=14,
    )
    img = sp.get_pixmap(dpi=150).tobytes("png")
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)
    page.insert_image(page.rect, stream=img)
    doc.set_metadata({"producer": "pine-demo", "creationDate": "D:20260101000000Z"})
    data = bytes(doc.tobytes(deflate=True))
    doc.close()
    return data


def _ground_truth() -> GroundTruth:
    def fact(
        metric: str,
        value: int | float,
        unit: str,
        ptype: str,
        start: str | None,
        end: str | None,
        source: str,
        currency: str | None = "USD",
        tol: float = 0.0,
    ) -> GroundTruthFact:
        return GroundTruthFact(
            metric=metric,
            value=value,
            unit=unit,
            currency=currency if unit == "currency" else None,
            period=Period(
                type=ptype,
                start=date.fromisoformat(start) if start else None,
                end=date.fromisoformat(end) if end else None,
            ),
            source_file=source,
            tolerance_pct=tol,
        )

    fy24 = ("fiscal_year", "2024-01-01", "2024-12-31")
    fy25 = ("fiscal_year", "2025-01-01", "2025-12-31")
    return GroundTruth(
        facts=[
            fact("arr", 12_000_000, "currency", "quarter", "2025-10-01",
                 "2025-12-31", "01_Northwind_SeriesB_Deck.pptx"),
            fact("revenue", 4_100_000, "currency", *fy24,
                 "02_Financials_FY2024_FY2025.xlsx"),
            fact("revenue", 9_700_000, "currency", *fy25,
                 "02_Financials_FY2024_FY2025.xlsx"),
            fact("cogs", 1_230_000, "currency", *fy24,
                 "02_Financials_FY2024_FY2025.xlsx"),
            fact("cogs", 2_522_000, "currency", *fy25,
                 "02_Financials_FY2024_FY2025.xlsx"),
            fact("cash_balance", 6_200_000, "currency", "point",
                 "2025-12-31", "2025-12-31", "02_Financials_FY2024_FY2025.xlsx"),
            fact("headcount", 84, "count", *fy25,
                 "02_Financials_FY2024_FY2025.xlsx", currency=None),
            fact("recurring_revenue", 10_200_000, "currency", *fy25,
                 "03_Customer_List.csv"),
            fact("bank_inflows", 9_900_000, "currency", *fy25,
                 "04_Bank_Statement_2025.csv", tol=1.0),
            fact("cash_balance", 6_200_000, "currency", "point",
                 "2025-12-31", "2025-12-31", "04_Bank_Statement_2025.csv"),
            fact("contract_value", 6_700_000, "currency", "custom",
                 "2025-01-01", "2027-12-31",
                 "05_Contracts/MSA_Acme_Corporation.pdf"),
            fact("net_burn", 450_000, "currency", "month", "2025-10-01",
                 "2025-10-31", "07_Board_Deck_Q4_2025.pdf"),
            fact("runway_months", 14, "months", "point", "2025-12-31",
                 "2025-12-31", "07_Board_Deck_Q4_2025.pdf", currency=None),
            fact("net_revenue_retention", 118, "percent", *fy25,
                 "07_Board_Deck_Q4_2025.pdf", currency=None),
            fact("top_customer_concentration", 22, "percent", *fy25,
                 "03_Customer_List.csv", currency=None, tol=1.0),
            fact("tam", 40_000_000_000, "currency", *fy25,
                 "01_Northwind_SeriesB_Deck.pptx"),
        ],
        contradictions=[
            GroundTruthContradiction(
                rule_id="R1",
                metric="arr",
                expected_values=[12_000_000, 9_700_000, 10_200_000],
                min_severity="high",
            ),
            GroundTruthContradiction(
                rule_id="R5",
                metric="fully_diluted_pct",
                expected_values=[103.0],
                min_severity="high",
            ),
        ],
        entities=[
            GroundTruthEntity(
                type="company",
                canonical_name="Northwind",
                aliases=["Northwind SaaS", "Northwind, Inc."],
            ),
            GroundTruthEntity(
                type="customer",
                canonical_name="Acme Corporation",
                aliases=["ACME Corporation", "Acme Corp."],
            ),
            GroundTruthEntity(
                type="customer", canonical_name="Helios Logistics"
            ),
            GroundTruthEntity(
                type="customer", canonical_name="Borealis Health"
            ),
            GroundTruthEntity(
                type="customer", canonical_name="Delta Freight"
            ),
            GroundTruthEntity(
                type="customer", canonical_name="Kestrel Bank"
            ),
            GroundTruthEntity(
                type="bank_account",
                canonical_name="Northwind Operating Account",
                aliases=["ACCT 4839201156"],
            ),
        ],
    )


# ---------------------------------------------------------------------------
# entry point


def build_demo(out_dir: str | Path, force: bool = False) -> Path:
    """Build the Northwind demo room into `out_dir`; returns the directory."""
    out = Path(out_dir)
    if out.exists() and any(out.iterdir()):
        if not force:
            raise FileExistsError(
                f"{out} already exists and is not empty (use --force to rebuild)"
            )
        shutil.rmtree(out)

    (out / "05_Contracts").mkdir(parents=True, exist_ok=True)
    (out / "08_Emails").mkdir(parents=True, exist_ok=True)

    (out / "01_Northwind_SeriesB_Deck.pptx").write_bytes(_deck_pptx())
    (out / "02_Financials_FY2024_FY2025.xlsx").write_bytes(_financials_xlsx())
    (out / "03_Customer_List.csv").write_bytes(_customer_list_csv())
    (out / "04_Bank_Statement_2025.csv").write_bytes(_bank_statement_csv())
    for name, data in _contracts().items():
        (out / "05_Contracts" / name).write_bytes(data)
    (out / "06_Cap_Table.xlsx").write_bytes(_cap_table_xlsx())
    (out / "07_Board_Deck_Q4_2025.pdf").write_bytes(_board_deck_pdf())
    for name, data in _emails().items():
        (out / "08_Emails" / name).write_bytes(data)
    (out / "09_Scanned_Bank_Letter.pdf").write_bytes(_scanned_letter())
    (out / "ground_truth.json").write_text(
        json.dumps(_ground_truth().model_dump(mode="json"), indent=2) + "\n"
    )
    return out
