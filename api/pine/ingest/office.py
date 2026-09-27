import io

from pine.ingest.types import (
    ParsedBlock,
    ParsedCell,
    ParsedPage,
    ParsedTable,
    ParseResult,
)
from pine.models.document import BlockKind


def parse_docx(data: bytes, filename: str) -> ParseResult:
    import docx

    doc = docx.Document(io.BytesIO(data))
    blocks: list[ParsedBlock] = []
    tables: list[ParsedTable] = []
    cursor = 0
    order = 0
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        kind = (
            BlockKind.heading
            if para.style and para.style.name.startswith("Heading")
            else BlockKind.paragraph
        )
        blocks.append(
            ParsedBlock(
                order=order, kind=kind, text=text,
                char_start=cursor, char_end=cursor + len(text),
            )
        )
        cursor += len(text) + 1
        order += 1
    for t_idx, table in enumerate(doc.tables):
        cells: list[ParsedCell] = []
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                cells.append(ParsedCell(row=r, col=c, text=cell.text))
        tables.append(
            ParsedTable(
                order=t_idx,
                n_rows=len(table.rows),
                n_cols=len(table.columns) if table.rows else 0,
                header_row=0,
                cells=cells,
            )
        )
    text = "\n".join(b.text for b in blocks)
    page = ParsedPage(
        page_no=1, width=0, height=0, text=text, blocks=blocks, tables=tables
    )
    meta = {"title": doc.core_properties.title, "author": doc.core_properties.author}
    return ParseResult(pages=[page], meta=meta)


def parse_pptx(data: bytes, filename: str) -> ParseResult:
    from pptx import Presentation

    pres = Presentation(io.BytesIO(data))
    pages: list[ParsedPage] = []
    for i, slide in enumerate(pres.slides):
        blocks: list[ParsedBlock] = []
        tables: list[ParsedTable] = []
        order = 0
        cursor = 0
        for shape in slide.shapes:
            if shape.has_table:
                tbl = shape.table
                cells = [
                    ParsedCell(row=r, col=c, text=cell.text)
                    for r, row in enumerate(tbl.rows)
                    for c, cell in enumerate(row.cells)
                ]
                tables.append(
                    ParsedTable(
                        order=len(tables),
                        n_rows=len(tbl.rows),
                        n_cols=len(tbl.columns),
                        header_row=0,
                        cells=cells,
                    )
                )
                continue
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            kind = BlockKind.heading if shape == slide.shapes.title else BlockKind.paragraph
            blocks.append(
                ParsedBlock(
                    order=order, kind=kind, text=text,
                    char_start=cursor, char_end=cursor + len(text),
                )
            )
            cursor += len(text) + 1
            order += 1
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                blocks.append(
                    ParsedBlock(
                        order=order, kind=BlockKind.slide_note, text=notes,
                        char_start=cursor, char_end=cursor + len(notes),
                    )
                )
        text = "\n".join(b.text for b in blocks)
        pages.append(
            ParsedPage(
                page_no=i + 1,
                width=pres.slide_width / 914400 if pres.slide_width else 0,
                height=pres.slide_height / 914400 if pres.slide_height else 0,
                text=text,
                blocks=blocks,
                tables=tables,
            )
        )
    return ParseResult(pages=pages)
